#!/usr/bin/env python3
"""Build the study deck: environment, scenarios, how Claude Code is run, and every hypothesis
with its token results. Plain white layout, python-pptx."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = sys.argv[1]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

INK = RGBColor(0x20, 0x20, 0x20)
GREY = RGBColor(0x66, 0x66, 0x66)
LINE = RGBColor(0x1F, 0x4E, 0x79)
HEAD_BG = RGBColor(0xF2, 0xF2, 0xF2)
FONT = "Calibri"
MARGIN = Inches(0.6)
CONTENT_W = prs.slide_width - 2 * MARGIN

slide_number = [0]


def _text(frame, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, mono=False):
    p = frame.paragraphs[0] if frame.paragraphs and frame.paragraphs[0].text == "" else frame.add_paragraph()
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Menlo" if mono else FONT
    r.font.color.rgb = color
    return p


def new_slide(title, subtitle=None):
    slide = prs.slides.add_slide(BLANK)
    slide_number[0] += 1
    tb = slide.shapes.add_textbox(MARGIN, Inches(0.35), CONTENT_W, Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    _text(tf, title, 28, bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.05), CONTENT_W, Emu(19050))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    if subtitle:
        sb = slide.shapes.add_textbox(MARGIN, Inches(1.12), CONTENT_W, Inches(0.5))
        sb.text_frame.word_wrap = True
        _text(sb.text_frame, subtitle, 14, color=GREY)
    nb = slide.shapes.add_textbox(prs.slide_width - Inches(1.2), prs.slide_height - Inches(0.45), Inches(0.8), Inches(0.3))
    _text(nb.text_frame, str(slide_number[0]), 10, color=GREY, align=PP_ALIGN.RIGHT)
    return slide


def bullets(slide, items, top=1.65, left=None, width=None, size=16, height=None):
    left = MARGIN if left is None else left
    width = CONTENT_W if width is None else width
    height = Inches(7.5 - top - 0.6) if height is None else height
    tb = slide.shapes.add_textbox(left, Inches(top), width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        level = 0
        text = item
        if isinstance(item, tuple):
            level, text = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        bullet = "•  " if level == 0 else "–  "
        r = p.add_run()
        r.text = bullet + text
        r.font.size = Pt(size - 2 * level)
        r.font.name = FONT
        r.font.color.rgb = INK
        p.space_after = Pt(6)
    return tb


def para(slide, text, top, size=14, left=None, width=None, color=INK, mono=False, height=None):
    left = MARGIN if left is None else left
    width = CONTENT_W if width is None else width
    height = Inches(0.6) if height is None else height
    tb = slide.shapes.add_textbox(left, Inches(top), width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = "Menlo" if mono else FONT
        r.font.color.rgb = color
    return tb


def table(slide, rows, top=1.65, size=11, col_widths=None, left=None, width=None, row_h=0.32, bold_first_col=False):
    left = MARGIN if left is None else left
    width = CONTENT_W if width is None else width
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, left, Inches(top), width, Inches(row_h * n_rows))
    tbl = shape.table
    # plain look: no banding style
    tblPr = shape._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = int(width * w / total)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(value)
            run.font.size = Pt(size)
            run.font.name = FONT
            run.font.color.rgb = INK
            run.font.bold = (r == 0) or (bold_first_col and c == 0)
            if c > 0 and r > 0 and _numeric(value):
                p.alignment = PP_ALIGN.RIGHT
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEAD_BG if r == 0 else RGBColor(0xFF, 0xFF, 0xFF)
    return shape


def _numeric(v):
    s = str(v).replace(",", "").replace("k", "").replace("%", "").replace("+", "").replace("-", "").replace("/", "").replace(".", "").replace(" ", "")
    return s.isdigit() and str(v) not in ("-",)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


